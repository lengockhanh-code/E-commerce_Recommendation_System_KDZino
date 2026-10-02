"""Create editable MerRec PowerPoint from the reviewed outline and local figures."""
from pathlib import Path
import sys,json,csv,math
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/".cache/slide_tools"))
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from PIL import Image

OUT=Path(__file__).resolve().parent
DATA=json.loads((OUT/"slides.json").read_text(encoding="utf-8"))
IMG=OUT/"images"
prs=Presentation()
prs.slide_width=Inches(13.333333)
prs.slide_height=Inches(7.5)
prs.core_properties.title="MerRec — Hệ thống gợi ý sản phẩm"
prs.core_properties.subject="Phân tích dữ liệu, đặc trưng và thực nghiệm mô hình"
prs.core_properties.author="MerRec"
prs.core_properties.keywords="MerRec, EDA, recommendation, Top-K"
prs.core_properties.comments="38 slide chính + 2 phụ lục demo. SASRec, DCN, LightGBM chưa có kết quả."
C={"blue":"244BFF","navy":"10213C","ink":"172B49","muted":"5B6B83","pale":"F0F4FC","line":"DFE6F2","teal":"008E89","amber":"F2B544","white":"FFFFFF","darkpale":"203553"}
FONT="Arial"
def rgb(c):return RGBColor.from_string(C.get(c,c))
def box(s,x,y,w,h,fill="pale",line=None,radius=False):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
    if line:sh.line.color.rgb=rgb(line);sh.line.width=Pt(.8)
    else:sh.line.fill.background()
    if radius:
        try:sh.adjustments[0]=.09
        except Exception:pass
    return sh
def text(s,x,y,w,h,value,size=20,color="ink",bold=False,align=PP_ALIGN.LEFT):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=0
    tf.margin_top=tf.margin_bottom=0
    for k,line in enumerate(str(value).split("\n")):
        p=tf.paragraphs[0] if k==0 else tf.add_paragraph()
        p.text=line;p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(6)
        p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
    return sh
def label(s,x,y,w,t,color="blue",size=11):
    return text(s,x,y,w,.28,t,size,color,True)
def line(s,x1,y1,x2,y2,color="line",width=1.4):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    sh.line.color.rgb=rgb(color);sh.line.width=Pt(width)
    return sh
def circle(s,x,y,d,fill="blue"):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d))
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill);sh.line.fill.background();return sh
def arrow(s,x,y,w=.3,h=.2,color="blue"):
    sh=s.shapes.add_shape(MSO_SHAPE.CHEVRON,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(color);sh.line.fill.background()
def bullets(s,items,x,y,w,size=19,gap=.85,color="ink"):
    for i,b in enumerate(items):
        circle(s,x,y+i*gap+.11,.065,"blue")
        text(s,x+.21,y+i*gap,w-.21,gap-.09,b,size,color)
def callout(s,t,y=6.37,fill="pale",color="ink",h=.48):
    box(s,.62,y,12.1,h,fill,radius=True)
    text(s,.82,y+.09,11.7,h-.12,t,13,color)
def pill(s,x,y,w,t,fill="blue",color="white"):
    box(s,x,y,w,.36,fill,radius=True);text(s,x+.1,y+.07,w-.2,.2,t,10,color,True)
def metric(s,x,y,w,h,value,caption,detail="",dark=False):
    box(s,x,y,w,h,"darkpale" if dark else "pale",radius=True)
    text(s,x+.22,y+.2,w-.44,.68,value,35,"white" if dark else "blue",True)
    text(s,x+.22,y+.94,w-.44,.58,caption,18,"white" if dark else "ink",True)
    if detail:text(s,x+.22,y+1.57,w-.44,h-1.62,detail,13,"B6C6DD" if dark else "muted")
def picture(s,name,x,y,w,h):
    path=IMG/name
    with Image.open(path) as im:iw,ih=im.size
    r=min(w/iw,h/ih);ww=iw*r;hh=ih*r
    sh=s.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))
    sh._element.nvPicPr.cNvPr.set("descr",name)
    return sh
def table(s,x,y,w,h,heads,rows,widths=None,font=16):
    shape=s.shapes.add_table(len(rows)+1,len(heads),Inches(x),Inches(y),Inches(w),Inches(h))
    t=shape.table
    if widths:
        for col,cw in zip(t.columns,widths):col.width=Inches(cw)
    for i,row in enumerate([heads]+rows):
        for j,val in enumerate(row):
            c=t.cell(i,j);c.text=str(val);c.margin_left=Inches(.14);c.margin_right=Inches(.1)
            c.margin_top=Inches(.1);c.margin_bottom=Inches(.07);c.vertical_anchor=MSO_ANCHOR.MIDDLE
            c.fill.solid();c.fill.fore_color.rgb=rgb("blue" if i==0 else ("pale" if i%2 else "white"))
            for p in c.text_frame.paragraphs:
                p.font.name=FONT;p.font.size=Pt(font);p.font.color.rgb=rgb("white" if i==0 else "ink");p.font.bold=(i==0)
    return shape
def flow(s,steps,y=2.65,h=1.55,x=.65,w=12.0):
    gap=.32;cw=(w-gap*(len(steps)-1))/len(steps)
    for i,(title,sub) in enumerate(steps):
        xx=x+i*(cw+gap);box(s,xx,y,cw,h,"pale",radius=True)
        label(s,xx+.17,y+.15,cw-.3,f"{i+1:02d}")
        text(s,xx+.17,y+.49,cw-.34,.52,title,19,"ink",True)
        text(s,xx+.17,y+1.04,cw-.34,h-1.1,sub,13,"muted")
        if i<len(steps)-1:arrow(s,xx+cw+.045,y+h/2-.1,.23,.2)
def standard(n,section,title):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid();s.background.fill.fore_color.rgb=rgb("white")
    box(s,0,0,.12,7.5,"blue")
    label(s,.62,.27,10.9,section.upper(),size=10)
    text(s,.62,.78,12.1,.95,title,29 if len(title)<66 else 26,"ink",True)
    line(s,.62,7.02,12.72,7.02)
    text(s,.62,7.12,9.9,.2,"MERREC  /  PHÂN TÍCH DỮ LIỆU & HỆ THỐNG GỢI Ý",9,"muted")
    text(s,11.65,7.09,1.07,.25,f"{n:02d} / 40",11,"blue",True,PP_ALIGN.RIGHT)
    box(s,.62,6.99,12.1*n/40,.035,"blue")
    return s
def img_and_text(s,name,items,foot="",size=19):
    box(s,.62,1.88,7.65,4.24,"white","line",True)
    picture(s,name,.74,2.02,7.41,3.95)
    bullets(s,items,8.6,2.08,4.0,size,.97)
    if foot:callout(s,foot)
def notes(s,row):
    n,sec,title,bs,vis,imgs,sources,note=row
    detail=f"SLIDE {n:02d} — {title}\n\n"+"\n".join(bs)
    if note:detail+="\n\nGhi chú: "+note
    if sources:detail+="\n\nNguồn tham khảo (URL hoặc đường dẫn trong D:/MerRec):\n"+"\n".join(sources)
    if imgs:detail+="\n\nẢnh: "+", ".join(imgs)
    detail+="\n\nSASRec, DCN và LightGBM Ranker chưa có kết quả để báo cáo theo xác nhận người dùng."
    s.notes_slide.notes_text_frame.text=detail

for row in DATA:
    n,section,title,bs,visual,imgs,sources,note=row
    s=standard(n,section,title)
    if n==1:
        s.background.fill.fore_color.rgb=rgb("navy")
        # Cover all standard elements with a dark canvas.
        box(s,0,0,13.333333,7.5,"navy")
        pill(s,.72,.55,2.15,"RECOMMENDER SYSTEM")
        text(s,.72,1.55,7.2,1.15,"MerRec",64,"white",True)
        text(s,.76,2.72,7.4,1.32,"Gợi ý sản phẩm từ\ndữ liệu hành vi",35,"white",True)
        text(s,.78,4.39,6.65,.84,"Phân tích dữ liệu · Xây dựng đặc trưng\nThực nghiệm mô hình Top-K",20,"B6C6DD")
        text(s,.78,6.28,7.0,.6,"Nhóm thực hiện: [Tên nhóm]    |    GVHD: [Tên giảng viên]",12,"B6C6DD")
        for i,(v,t) in enumerate([("174,87M","tương tác"),("30,05M","item"),("05/2023","dữ liệu quan sát")]):
            yy=1.38+i*1.65;box(s,9.0,yy,3.4,1.35,"darkpale",radius=True)
            text(s,9.27,yy+.18,2.9,.54,v,30,"white",True)
            text(s,9.29,yy+.85,2.85,.28,t,14,"B6C6DD")
        label(s,.78,7.06,6.0,"38 SLIDE NGHIÊN CỨU + 2 TRANG DEMO RIÊNG","B6C6DD",10)
    elif n==2:
        entries=[("01","Introduction","Bối cảnh và mục tiêu"),("02","Problem Definition","Bài toán và thách thức"),("03","Dataset Description","Quy mô và cấu trúc"),("04","Data Preprocessing","Làm sạch và chia tập"),("05","EDA","Khám phá dữ liệu"),("06","Feature Engineering","Biểu diễn dữ liệu"),("07","Models","Các hướng mô hình"),("08","Model Training","Huấn luyện"),("09","Results","Kết quả và giao thức"),("11","Conclusion","Kết luận và hướng đi")]
        for i,(no,t,sub) in enumerate(entries):
            col=i//5;rr=i%5;x=.73+col*6.1;y=1.94+rr*.8
            circle(s,x,y,.43);text(s,x,y+.11,.43,.23,no,12,"white",True,PP_ALIGN.CENTER)
            text(s,x+.63,y-.01,4.9,.3,t,18,"ink",True);text(s,x+.63,y+.35,4.9,.24,sub,12,"muted")
        callout(s,"10  Web Application / Demo được tách thành phụ lục, trang 39–40.")
    elif n==3:
        metric(s,.72,1.98,5.82,2.02,">25 tỷ USD","Quy mô TMĐT Việt Nam · 2024")
        metric(s,6.8,1.98,5.82,2.02,"+20%","Tăng trưởng so với năm 2023")
        bullets(s,["Nhu cầu ứng dụng: giúp người mua khám phá sản phẩm phù hợp trong catalog đa dạng.",
                   "Bài toán nghiên cứu: khai thác dữ liệu hành vi và metadata cho gợi ý C2C."],.88,4.37,11.6,20,.78)
        callout(s,"Bối cảnh: Việt Nam. Dữ liệu thực nghiệm: MerRec từ Mercari — hai phạm vi khác nhau.",y=6.16,h=.43)
        source=text(s,.85,6.69,11.7,.2,"Nguồn: Bộ Công Thương, báo cáo năm 2024 (18/02/2025); Li et al., MerRec (2024).",9,"muted")
        source.text_frame.paragraphs[0].runs[0].hyperlink.address=sources[0]
    elif n==4:
        box(s,.73,1.98,5.8,3.97,"pale",radius=True)
        label(s,1.0,2.26,5.2,"GÓC NHÌN NGƯỜI DÙNG",size=13)
        bullets(s,["Cần thu hẹp lựa chọn theo mối quan tâm.",
                   "Từ khóa cần diễn đạt rõ nhu cầu.",
                   "Danh sách phổ biến không tự phản ánh sở thích riêng."],1.0,2.94,5.13,20,.86)
        box(s,6.8,1.98,5.8,3.97,"pale",radius=True)
        label(s,7.08,2.26,5.2,"BẰNG CHỨNG TỪ MERREC",size=13)
        for j,(v,t) in enumerate([("30,05M","item trong lát dữ liệu"),("71,17%","item có dưới 5 tương tác"),("84,55%","sự kiện là lượt xem")]):
            yy=2.83+j*.94
            text(s,7.08,yy,2.16,.52,v,29,"blue",True)
            text(s,9.32,yy+.11,2.85,.5,t,17)
        callout(s,"Động lực: tận dụng nhiều tín hiệu để gợi ý một danh sách ngắn phù hợp hơn.",y=6.27,h=.53)
    elif n==5:
        reasons=[("01 · Tính thực tiễn","Hỗ trợ khám phá sản phẩm từ nhu cầu và lịch sử người dùng."),
                 ("02 · Tính học thuật","Nghiên cứu implicit feedback, long-tail, cold-start và chuỗi."),
                 ("03 · Dữ liệu phù hợp","MerRec có timestamp, nhiều hành vi và metadata để thực nghiệm."),
                 ("04 · Khả năng ứng dụng","Kết nối xử lý dữ liệu, mô hình và giao diện minh họa.")]
        for i,(t,d) in enumerate(reasons):
            x=.73+(i%2)*6.08;y=1.98+(i//2)*2.07
            box(s,x,y,5.81,1.79,"pale",radius=True)
            text(s,x+.24,y+.23,5.28,.43,t,22,"blue",True)
            text(s,x+.24,y+.86,5.28,.65,d,19)
        callout(s,"Giá trị kỳ vọng của đề tài; chưa phải bằng chứng tăng doanh thu hoặc hiệu quả online.",y=6.33,h=.48)
    elif n==6:
        box(s,.73,1.96,11.87,1.04,"navy",radius=True)
        text(s,.98,2.19,11.34,.64,"Xây dựng và đánh giá gợi ý Top-K từ dữ liệu hành vi MerRec",25,"white",True,PP_ALIGN.CENTER)
        columns=[("CÔNG VIỆC","Xử lý dữ liệu và EDA.\nTạo đặc trưng.\nThực nghiệm mô hình."),
                 ("Ý NGHĨA","Hiểu sparsity và cold-start.\nĐọc đúng giao thức đánh giá.\nLàm cơ sở tích hợp."),
                 ("PHẠM VI","Dữ liệu tháng 05/2023.\nĐánh giá offline.\nDemo web tách riêng.")]
        for i,(t,d) in enumerate(columns):
            x=.73+i*4.08
            box(s,x,3.43,3.71,2.36,"pale",radius=True)
            label(s,x+.22,3.7,3.27,t,size=13)
            text(s,x+.22,4.25,3.27,1.43,d,18)
        callout(s,"SASRec, DCN và LightGBM chưa hoàn tất; chưa có A/B test hoặc kiểm chứng kinh doanh.",y=6.27,h=.52)
    elif n==7:
        flow(s,[("Đầu vào","Lịch sử + hành vi + metadata"),("Mô hình","Retrieval / next-item"),("Top-K","Danh sách item xếp hạng")],y=1.99,h=1.74)
        for x,v,t in [(.89,"99,999843%","Độ thưa user–item"),(7.07,"32,75%","Dòng TEST chứa cold item")]:
            text(s,x,4.19,5.18,.62,v,34,"blue",True)
            text(s,x,4.99,5.18,.42,t,19,"muted")
        callout(s,"Recall@K · HitRate@K · NDCG@K — đọc cùng coverage, tập user và candidate universe.",y=6.12,h=.66)
    elif n==8:
        for x,v,t in [(.74,"174,87M","tương tác"),(4.83,"2,77M","người dùng"),(8.92,"30,05M","item_id duy nhất")]:
            box(s,x,1.94,3.69,1.29,"pale",radius=True)
            text(s,x+.2,2.13,3.27,.56,v,31,"blue",True)
            text(s,x+.2,2.8,3.27,.26,t,15)
        table(s,.74,3.57,11.88,2.18,["Nhóm trường","Thông tin chính"],
              [["Định danh","user_id · item_id · product_id"],
               ["Hành vi / thời gian","event_id · stime · session_id · sequence_id"],
               ["Metadata","name · price · category · brand · condition"]],[3.29,8.59],17)
        callout(s,"312 file · ~21,57 GB · 913.718 product_id · Đồ án dùng lát tháng 05/2023.",y=6.03,h=.48)
        source=text(s,.85,6.64,11.7,.25,"Nguồn: Li et al., MerRec (2024) và bảng tổng hợp local. Dataset gốc 6 tháng; đồ án dùng 1 tháng.",10,"muted")
        source.text_frame.paragraphs[0].runs[0].hyperlink.address=sources[0]
    elif n==9:
        table(s,.67,1.94,12,4.53,["Tên sự kiện gốc","Tên chuẩn","Ý nghĩa"],[
        ["item_view","view","Xem sản phẩm"],["item_like","like","Thích sản phẩm"],["item_add_to_cart_tap","cart","Thêm vào giỏ"],["offer_make","offer","Đưa ra đề nghị"],["buy_start","buy_start","Bắt đầu mua"],["buy_comp","buy_comp","Hoàn tất mua"]],[5.0,2.3,4.7],18)
    elif n==10:
        stages=[("Parquet → VIEW","Đọc bằng DuckDB"),("Chuẩn hóa","ID · thời gian · metadata"),("Làm sạch","Lọc lỗi và khử trùng"),("Temporal split","TRAIN → VAL → TEST"),("Đặc trưng","ID maps · CF pairs · catalog"),("Đầu ra","Snapshots · cold-start cohorts")]
        for i,(t,d) in enumerate(stages):
            x=.68+(i%3)*4.14;y=2.0+(i//3)*2.04
            box(s,x,y,3.85,1.76,"pale",radius=True);label(s,x+.2,y+.17,3.3,f"BƯỚC {i+1:02d}")
            text(s,x+.2,y+.6,3.45,.43,t,20,"ink",True);text(s,x+.2,y+1.19,3.45,.34,d,14,"muted")
        callout(s,"Không nạp toàn bộ raw vào Pandas; xử lý theo VIEW, batch và file Parquet.")
    elif n==11:
        metric(s,.68,1.99,5.2,1.94,"174.872.167","Dòng trước xử lý")
        arrow(s,6.2,2.72,.7,.5)
        metric(s,7.4,1.99,5.2,1.94,"174.725.659","Dòng sau làm sạch")
        label(s,.8,4.25,11.8,"KHÓA KHỬ TRÙNG")
        box(s,.72,4.68,11.89,.66,"blue",radius=True)
        text(s,.96,4.87,11.4,.39,"user · item · event · timestamp · session · sequence",21,"white",True,PP_ALIGN.CENTER)
        text(s,.8,5.62,11.6,.54,"Tổng giảm 146.508 dòng; giữ một sự kiện trên mỗi khóa sau bước lọc hợp lệ.",20)
        callout(s,"Chuẩn hóa kiểu ID, trim chuỗi và parse timestamp theo đúng kiểu dữ liệu nguồn.")
    elif n==12:
        img_and_text(s,imgs[0],["Color thiếu 95,61%; size_name thiếu 62,74%.","Brand_name thiếu 19,20%; các trường tùy chọn không đồng nghĩa sự kiện vô dụng.","Danh mục dùng cấp cha; brand/condition/shipper dùng __UNK__.","Color không xuất vào interactions train/val/test."],"Tỷ lệ missing theo dòng tương tác; giữ NULL của size ở bước chuẩn hóa.",18)
    elif n==13:
        picture(s,imgs[0],.75,1.85,11.8,2.28)
        table(s,.72,4.3,11.9,1.62,["Tập","Thời gian UTC","Số tương tác"],[["TRAIN","01–25/05","145.998.253"],["VAL","26–28/05","16.885.121"],["TEST","29/05–31/05 00:00","11.842.285"]],[2.1,5.7,4.1],16)
        callout(s,"Integrity đã lưu: TRAIN kết thúc trước VAL; VAL kết thúc trước TEST.")
    elif n==14:
        img_and_text(s,imgs[0],["Mapping chỉ xây từ TRAIN.","2.581.378 user và 27.328.461 item.","TEST cold-item: 32,75% dòng.","Tách warm, cold-user, cold-item và cold-both."],"Cold-user và cold-item có phần giao nhau; không cộng ba cột thành tổng.",19)
    elif n==15:
        picture(s,imgs[0],.85,1.76,11.66,4.68)
        callout(s,"Tỷ lệ đếm hành vi, không phải conversion rate. Cart và offer đã dùng mapping đúng.",y=6.45,h=.4)
    elif n==16:
        img_and_text(s,imgs[0],["Mean 63,23; median chỉ 12 tương tác/user.","P95 = 275; P99 = 791.","13,31% user có đúng một tương tác.","31,36% user có dưới 5 tương tác."],"Histogram trên mẫu 100.000 user; các số thống kê từ toàn bộ user.",19)
    elif n==17:
        picture(s,imgs[0],.7,1.91,5.88,3.74);picture(s,imgs[1],6.82,1.91,5.82,3.74)
        for x,v,t in [( .84,"2","Trung vị tương tác/item"),(4.93,"71,17%","Item có dưới 5 tương tác"),(9.04,"85,59%","Item có dưới 10 tương tác")]:
            text(s,x,5.75,3.6,.5,v,30,"blue",True);text(s,x,6.34,3.6,.34,t,14,"muted")
    elif n==18:
        picture(s,imgs[0],.74,1.94,5.78,3.44);picture(s,imgs[1],6.84,1.94,5.78,3.44)
        for x,v,t in [(.84,"34,83M","Session theo user + session_id"),(4.93,"2","Trung vị tương tác/session"),(9.04,"37 giây","Trung vị thời lượng")]:
            text(s,x,5.66,3.6,.52,v,30,"blue",True);text(s,x,6.26,3.6,.43,t,14,"muted")
    elif n==19:
        picture(s,imgs[0],.72,1.92,11.9,3.66)
        box(s,.73,5.82,11.85,.79,"pale",radius=True)
        text(s,.95,5.98,2.16,.4,"31/05: 88",22,"blue",True)
        text(s,3.18,5.98,9.0,.42,"Ngày cuối không đầy đủ; không diễn giải là nhu cầu thị trường giảm.",18)
    elif n==20:
        img_and_text(s,imgs[0],["Women: 59,15 triệu tương tác, khoảng 33,83%.","Toys & Collectibles: 32,90 triệu.","Men: 15,67 triệu.","Category phân cấp hỗ trợ biểu diễn nội dung và retrieval."],"Trục số liệu là tương tác, không phải số item hoặc số người mua.",18)
    elif n==21:
        img_and_text(s,imgs[0],["Mean 62,69; median 25.","P95 = 225; P99 = 636.","Giá quan sát từ 1 đến 5.000.","Two-Tower dùng log(1+price), sau đó chuẩn hóa."],"Theo đơn vị gốc của dataset; số liệu tổng hợp trên dòng tương tác.",19)
    elif n==22:
        table(s,.72,1.98,11.9,3.79,["Phát hiện","Quyết định / hướng xử lý"],[
        ["Sparsity ≈ 99,999843%","Thử implicit CF và biểu diễn metadata"],
        ["25,53% tương tác là lặp thêm","Tổng hợp tín hiệu theo cặp user–item"],
        ["Nhiều user/item ít lịch sử","Phân tích cold-start và phương án fallback"],
        ["One-pass 5-core loại 22,33%","Chỉ là chẩn đoán, chưa áp dụng như bộ lọc cuối"]
        ],[5.4,6.5],18)
        callout(s,"Không loại bỏ hàng loạt long-tail trước khi đánh giá ảnh hưởng tới coverage.",y=6.16,h=.59)
    elif n==23:
        img_and_text(s,imgs[0],["implicit_score(u,i) = tổng trọng số sự kiện.","Giữ thêm số sự kiện theo loại.","Giữ thời điểm tương tác cuối.","109.718.519 cặp user–item trong TRAIN."],"Trọng số do dự án đặt; không phải tham số mô hình tự học.",19)
    elif n==24:
        for x,head,items in [( .72,"CONTENT-BASED",["Name · category · brand · condition","Hashing TF-IDF","262.144 chiều","Unigram + bigram"]),(6.8,"TWO-TOWER",["7 trường metadata + price","Hashed embeddings","Log-price chuẩn hóa","Ghép vector → MLP"])]:
            box(s,x,1.97,5.81,3.96,"pale",radius=True);label(s,x+.26,2.25,5.2,head,size=14)
            bullets(s,items,x+.28,3.0,5.19,20,.65)
        callout(s,"Hashing và xử lý theo batch giúp kiểm soát vocabulary và bộ nhớ.")
    elif n==25:
        flow(s,[("i₁ / view","Item + event"),("i₂ / like","Item + event"),("i₃ / cart","Item + event"),("i₄ ?","Target tiếp theo")],y=2.26,h=1.73)
        for x,v,t in [( .83,"32 chiều","Item embedding"),(4.95,"16 chiều","Event embedding"),(9.02,"100 bước","Giới hạn chuỗi")]:
            text(s,x,4.6,3.6,.62,v,31,"blue",True);text(s,x,5.29,3.6,.44,t,17,"muted")
        callout(s,"PAD = 0 · OOV = 1 · Hành vi là feature, nhãn là item tiếp theo.")
    elif n==26:
        categories=[("Baseline / fallback","Popularity · Trending","Đã có snapshot"),("Nội dung / đồng xuất hiện","Content-Based · Co-visitation","Đã có kết quả TEST"),("Học biểu diễn","ALS · GRU4Rec · Two-Tower","Metric theo từng giao thức")]
        for i,(t,m,st) in enumerate(categories):
            x=.71+i*4.12;box(s,x,2.06,3.86,3.0,"pale",radius=True)
            label(s,x+.21,2.31,3.4,f"NHÓM {i+1:02d}")
            text(s,x+.21,2.91,3.43,.76,t,23,"ink",True)
            text(s,x+.21,3.89,3.43,.58,m,17)
            text(s,x+.21,4.67,3.43,.28,st,12,"teal")
        callout(s,"Chưa hoàn tất: SASRec · DCN · LightGBM Ranker — không có số liệu trong bảng kết quả.",y=5.74,h=.83)
    elif n==27:
        box(s,.74,2.04,11.83,1.23,"navy",radius=True)
        text(s,.96,2.42,11.36,.55,"Trending = Σ weight × 0,5 ^ (tuổi sự kiện / 3 ngày)",27,"white",True,PP_ALIGN.CENTER)
        for i,(day,v) in enumerate([("0 ngày","100%"),("3 ngày","50%"),("6 ngày","25%"),("9 ngày","12,5%")]):
            x=.87+i*3.08;text(s,x,3.97,2.78,.58,v,33,"blue",True)
            text(s,x,4.72,2.78,.34,day,17,"muted")
        text(s,.85,5.42,11.65,.56,"Minh họa hệ số time-decay; không phải số liệu đo từ một mô hình.",17,"muted")
        callout(s,"Snapshot TRAIN dùng offline; snapshot toàn bộ processed phục vụ catalog/demo.")
    elif n==28:
        flow(s,[("Metadata","TF-IDF sản phẩm"),("Hồ sơ user","≤100 item lịch sử"),("So khớp","Vector chuẩn hóa"),("Top-K","Lọc item đã xem")],y=2.18,h=1.75)
        bullets(s,["Recency decay = 0,97 để tăng vai trò lịch sử gần.","VAL: lịch sử TRAIN. TEST: lịch sử TRAIN + VAL."],.83,4.42,11.7,21,.75)
        callout(s,"Metadata và IDF dùng full catalog; cần ghi rõ giới hạn của temporal evaluation.",y=6.22,h=.57)
    elif n==29:
        flow(s,[("Session","Bucket tương tác"),("Sliding window","Tạo cặp item"),("Top-K láng giềng","Reduce / checkpoint"),("Gợi ý","Tổng hợp từ lịch sử")],y=2.12,h=1.72)
        table(s,.76,4.27,11.81,1.57,["History","Decay","Reverse weight","Popularity penalty","Window"],[["30 item","0,9","0,4","0,1","5"]],[2.25,1.8,2.6,3.15,2.01],16)
        callout(s,"Chọn cấu hình trên VAL; refit TRAIN + VAL trước đánh giá TEST.")
    elif n==30:
        # Editable conceptual matrices; no synthetic result values.
        for gx,gy,rr,cc,cell,col in [(1.0,2.12,4,6,.39,"blue"),(5.25,2.12,4,2,.39,"teal"),(8.15,2.12,2,6,.39,"blue")]:
            for a in range(rr):
                for b in range(cc):
                    fill=col if (a+b)%3 else "line"
                    box(s,gx+b*(cell+.055),gy+a*(cell+.055),cell,cell,fill)
        text(s,3.97,2.65,.8,.74,"≈",40,"muted",True)
        text(s,6.72,2.63,.7,.6,"×",36,"muted",True)
        text(s,.99,4.12,3.6,.5,"Ma trận tương tác",20,"ink",True)
        text(s,5.05,4.12,2.7,.5,"User factors",20,"ink",True)
        text(s,8.16,4.12,3.7,.5,"Item factors",20,"ink",True)
        table(s,.77,5.03,11.8,1.11,["Factors","Regularization","Alpha","Iterations"],[["32","0,15","40","12"]],None,17)
        callout(s,"Sơ đồ nguyên lý; cấu hình từ ALS artifact. Chọn model bằng POSITIVE NDCG@20.",y=6.42,h=.4)
    elif n==31:
        flow(s,[("Item + event","32 + 16 chiều"),("GRU × 2","Hidden = 128"),("Ngữ cảnh","Vector người dùng"),("Next item","Sampled softmax")],y=2.18,h=1.76)
        bullets(s,["Item embedding trên CPU; GRU và dense được thiết kế chạy GPU.","64 negatives khi train; catalog lớn được truy cập theo batch."],.84,4.48,11.7,20,.7)
        pill(s,.87,6.08,3.18,"METRIC HIỆN CÓ: SAMPLED VAL")
        text(s,4.3,6.14,8.12,.43,"Không gọi là full-catalog TEST.",20,"muted")
    elif n==32:
        for x,head,body in [(1.0,"USER TOWER","User ID\n↓\nEmbedding 64 chiều"),(7.18,"ITEM TOWER","7 metadata × 16 + price_z\n↓\nMLP 256 → 64 chiều")]:
            box(s,x,1.98,5.12,2.62,"pale",radius=True);label(s,x+.24,2.24,4.5,head,size=14)
            text(s,x+.24,2.85,4.65,1.5,body,21,"ink",True,PP_ALIGN.CENTER)
        line(s,3.56,4.64,6.68,5.09,"blue",2);line(s,9.73,4.64,6.68,5.09,"blue",2)
        box(s,3.34,5.08,6.67,.69,"blue",radius=True)
        text(s,3.56,5.27,6.2,.36,"Chuẩn hóa vector → tích vô hướng",22,"white",True,PP_ALIGN.CENTER)
        callout(s,"Symmetric InfoNCE + in-batch negatives · FAISS chứa 27.328.461 TRAIN item.",y=6.31,h=.49)
    elif n==33:
        table(s,.71,1.98,11.92,3.71,["Mô hình","Quy mô / batch","Vòng đã lưu","Chọn checkpoint"],[
        ["ALS","32 factors","12 iterations","POSITIVE NDCG@20"],
        ["GRU4Rec","Batch 128; ≤1M windows/epoch","8 epoch","Sampled VAL Recall@20"],
        ["Two-Tower","Batch 2.048","15 epoch","VAL POSITIVE NDCG@20"]
        ],[2.0,3.62,2.05,4.25],17)
        callout(s,"Đọc batch/row-group · Lưu checkpoint · Resume khi phiên chạy bị gián đoạn.",y=6.03,h=.71)
    elif n==34:
        picture(s,imgs[0],.72,1.86,11.87,3.77)
        text(s,.88,5.89,5.65,.38,"GRU4Rec: best checkpoint tại epoch 6",19,"blue",True)
        text(s,6.94,5.89,5.46,.38,"Two-Tower: best tại epoch 15",19,"teal",True)
        callout(s,"Hai hàm mục tiêu khác nhau: chỉ đọc xu hướng loss riêng từng mô hình.",y=6.46,h=.38)
    elif n==35:
        picture(s,imgs[0],.72,1.94,7.15,4.16)
        bullets(s,["Recall: tỷ lệ target tìm thấy.","HitRate: user có ít nhất một hit.","NDCG: tính thêm vị trí của hit.","Luôn ghi split, user, candidates và cách lọc."],8.3,2.17,4.16,19,.88)
        callout(s,"Sampled VAL, candidate pool và full-index là các giao thức khác nhau.",y=6.4,h=.44)
    elif n==36:
        result=list(csv.DictReader((OUT/"ket_qua_trich_xuat.csv").open(encoding="utf-8-sig")))
        table(s,.71,2.01,11.92,3.17,["Mô hình","Recall@20","NDCG@20","HitRate@20"],[
        [r["model"],f'{float(r["Recall@20"])*100:.4f}%',f'{float(r["NDCG@20"]):.6f}',f'{float(r["HitRate@20"])*100:.4f}%'] for r in result
        ],[3.82,2.7,2.7,2.7],20)
        pill(s,.81,5.51,2.43,"TEST · TARGET POSITIVE")
        text(s,3.51,5.48,9.01,.67,"Khác tập user, candidate universe và coverage; chưa dùng để xếp hạng mô hình.",18)
        callout(s,"SASRec / DCN / LightGBM chưa có kết quả; GRU4Rec sampled VAL trình bày riêng.",y=6.33,h=.53)
    elif n==37:
        picture(s,imgs[0],.7,1.92,8.05,4.32)
        metric(s,9.04,2.14,3.58,2.12,"15,69%","Recall@20 · epoch 6")
        text(s,9.17,4.61,3.2,1.3,"1.810 user\n1 target + 255 negatives\nSampled validation",18,"muted")
        callout(s,"Bước tiếp theo: thống nhất cohort/candidates và đo full-index trước khi chọn model.")
    elif n==38:
        for x,head,items in [( .72,"ĐÃ THỰC HIỆN",["EDA và preprocessing 174,87M sự kiện.","Nhiều hướng mô hình và artifact đã lưu."]),(4.84,"ĐIỀU RÚT RA",["Long-tail và cold-item ảnh hưởng coverage.","Giao thức đánh giá quyết định cách đọc metric."]),(8.96,"BƯỚC TIẾP THEO",["Thống nhất benchmark; hoàn thiện SASRec/DCN/LightGBM.","Tích hợp và đo chất lượng phục vụ."])]:
            box(s,x,2.04,3.87,3.66,"pale",radius=True);label(s,x+.21,2.31,3.4,head,size=13)
            bullets(s,items,x+.22,3.02,3.39,19,1.2)
        callout(s,"Kết thúc nội dung nghiên cứu · Demo nằm riêng trong hai trang tiếp theo.",y=6.13,h=.63)
    elif n==39:
        label(s,.75,1.75,11.9,"PHỤ LỤC DEMO · 30–60 GIÂY",size=12)
        ph=box(s,.75,2.24,8.44,4.27,"pale","line",True)
        text(s,1.15,3.65,7.64,.65,"CHÈN ẢNH TRANG CHỦ MERREC",24,"muted",True,PP_ALIGN.CENTER)
        text(s,1.35,4.46,7.24,.45,"Khung ảnh để thay bằng screenshot thực tế",16,"muted",False,PP_ALIGN.CENTER)
        bullets(s,["Khám phá sản phẩm","Mở trang chi tiết","Xem khu vực gợi ý"],9.58,2.89,2.98,21,.97)
    elif n==40:
        for x,head in [(.76,"ẢNH CHI TIẾT SẢN PHẨM"),(6.84,"ẢNH DEMO TÙY CHỌN")]:
            box(s,x,2.03,5.74,3.18,"pale","line",True)
            text(s,x+.34,3.15,5.06,.8,head,21,"muted",True,PP_ALIGN.CENTER)
        text(s,.82,5.62,11.7,.66,"Cảm ơn thầy cô và các bạn",31,"blue",True,PP_ALIGN.CENTER)
        text(s,.82,6.33,11.7,.38,"Câu hỏi & thảo luận",20,"muted",False,PP_ALIGN.CENTER)
    notes(s,row)

dest=OUT/"MERREC_THUYET_TRINH_40_SLIDE.pptx"
prs.save(dest)
# Verify round trip, expected count, notes and shape boundaries.
verify=Presentation(dest)
assert len(verify.slides)==40
issues=[]
for idx,s in enumerate(verify.slides,1):
    assert s.has_notes_slide
    for sh in s.shapes:
        if sh.left < -100 or sh.top < -100 or sh.left+sh.width>verify.slide_width+200 or sh.top+sh.height>verify.slide_height+200:
            issues.append({"slide":idx,"shape":sh.name,"issue":"bounds"})
assert not issues,issues
(OUT/"POWERPOINT_KIEM_TRA.json").write_text(json.dumps({"slides":40,"editable_text_tables_diagrams":True,"bounds_issues":issues,"images_embedded":sum(1 for s in verify.slides for sh in s.shapes if sh.shape_type==13),"sources_in_speaker_notes":True},indent=2),encoding="utf-8")
print(json.dumps({"path":str(dest),"bytes":dest.stat().st_size,"slides":40,"bounds_issues":len(issues)},ensure_ascii=False))
